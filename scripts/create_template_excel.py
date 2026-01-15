"""
Script para crear un template Excel para carga masiva de productos
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows


def create_product_template():
    """Crea un archivo Excel template para carga masiva de productos"""
    
    # Crear workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    
    # Definir encabezados con descripciones
    headers = [
        "serial_number",
        "name", 
        "description",
        "brand_id",
        "category_id",
        "warranty_time",
        "warranty_unit",
        "cost",
        "retail_price",
        "status",
        "variant_color",
        "variant_size",
        "variant_size_unit",
        "variant_unit",
        "variant_branch_id",
        "variant_stock",
        "variant_min_stock",
        "image_paths"
    ]
    
    descriptions = [
        "Número de serie único del producto (requerido)",
        "Nombre del producto (requerido)",
        "Descripción detallada (opcional)",
        "ID de la marca (opcional, número)",
        "ID de la categoría (opcional, número)",
        "Tiempo de garantía (opcional, número)",
        "Unidad de garantía: DAYS, MONTHS, YEARS (opcional)",
        "Costo del producto (requerido, número decimal)",
        "Precio de venta (requerido, número decimal)",
        "Estado: ACTIVE, INACTIVE, DISCONTINUED (opcional, default: ACTIVE)",
        "Color de la variante: ROJO, AZUL, VERDE, AMARILLO, NARANJA, VIOLETA, ROSADO, MARRON, GRIS, BLANCO, NEGRO, BORDO (opcional)",
        "Talla/Tamaño de la variante (opcional, texto)",
        "Tipo de tamaño: CLOTHING, DIMENSIONS, WEIGHT, OTHER (opcional)",
        "Unidad: KG, G, LB, CM, M, INCH, XS, S, L, XL, XXL (opcional)",
        "ID de la sucursal (requerido para variante, número)",
        "Stock disponible (opcional, número, default: 0)",
        "Stock mínimo (opcional, número, default: 5)",
        "Rutas de imágenes separadas por ; (ejemplo: C:\\imgs\\img1.jpg;C:\\imgs\\img2.jpg)"
    ]
    
    # Escribir encabezados
    ws.append(headers)
    
    # Estilo para encabezados
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Escribir descripciones
    ws.append(descriptions)
    
    # Estilo para descripciones
    desc_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    desc_font = Font(italic=True, size=9)
    
    for cell in ws[2]:
        cell.fill = desc_fill
        cell.font = desc_font
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    # Agregar ejemplos
    examples = [
        [
            "BAT-IP12-001",
            "Batería iPhone 12",
            "Batería de alta capacidad para iPhone 12",
            "1",
            "1",
            "6",
            "MONTHS",
            "25000",
            "45000",
            "ACTIVE",
            "NEGRO",
            "",
            "",
            "",
            "1",
            "50",
            "5",
            "C:\\imgs\\bat-ip12-1.jpg;C:\\imgs\\bat-ip12-2.jpg"
        ],
        [
            "CASE-IP13-001",
            "Funda iPhone 13 Pro",
            "Funda de silicona premium",
            "2",
            "2",
            "",
            "",
            "8000",
            "15000",
            "ACTIVE",
            "AZUL",
            "",
            "",
            "",
            "1",
            "100",
            "10",
            "C:\\imgs\\case-ip13-azul.jpg"
        ]
    ]
    
    for example in examples:
        ws.append(example)
    
    # Ajustar anchos de columnas
    column_widths = {
        'A': 20,  # serial_number
        'B': 30,  # name
        'C': 40,  # description
        'D': 12,  # brand_id
        'E': 12,  # category_id
        'F': 15,  # warranty_time
        'G': 15,  # warranty_unit
        'H': 12,  # cost
        'I': 12,  # retail_price
        'J': 15,  # status
        'K': 15,  # variant_color
        'L': 15,  # variant_size
        'M': 18,  # variant_size_unit
        'N': 15,  # variant_unit
        'O': 18,  # variant_branch_id
        'P': 15,  # variant_stock
        'Q': 18,  # variant_min_stock
        'R': 50,  # image_paths
    }
    
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    # Establecer altura de filas
    ws.row_dimensions[1].height = 40
    ws.row_dimensions[2].height = 60
    
    # Crear hoja de referencia
    ws_ref = wb.create_sheet("Referencia")
    
    reference_data = {
        "Unidades de Garantía": ["DAYS", "MONTHS", "YEARS"],
        "Estados de Producto": ["ACTIVE", "INACTIVE", "DISCONTINUED"],
        "Colores": ["ROJO", "AZUL", "VERDE", "AMARILLO", "NARANJA", "VIOLETA", 
                   "ROSADO", "MARRON", "GRIS", "BLANCO", "NEGRO", "BORDO"],
        "Unidades de Tamaño": ["CLOTHING", "DIMENSIONS", "WEIGHT", "OTHER"],
        "Unidades": ["KG", "G", "LB", "CM", "M", "INCH", "XS", "S", "L", "XL", "XXL"]
    }
    
    col_idx = 1
    for title, values in reference_data.items():
        ws_ref.cell(row=1, column=col_idx, value=title)
        ws_ref.cell(row=1, column=col_idx).font = Font(bold=True)
        ws_ref.cell(row=1, column=col_idx).fill = PatternFill(
            start_color="70AD47", end_color="70AD47", fill_type="solid"
        )
        ws_ref.cell(row=1, column=col_idx).font = Font(color="FFFFFF", bold=True)
        
        for idx, value in enumerate(values, start=2):
            ws_ref.cell(row=idx, column=col_idx, value=value)
        
        ws_ref.column_dimensions[ws_ref.cell(row=1, column=col_idx).column_letter].width = 25
        col_idx += 1
    
    # Guardar archivo
    wb.save("template_carga_productos.xlsx")
    print("✓ Template Excel creado exitosamente: template_carga_productos.xlsx")


if __name__ == "__main__":
    create_product_template()
