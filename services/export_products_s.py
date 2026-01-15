"""
Servicio para exportar productos a Excel
"""
import os
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlmodel import Session, select
from sqlalchemy.orm import selectinload
from database.models.product import Product, ProductVariant


def export_products_to_excel(session: Session, filename: str = "productos_exportados.xlsx") -> str:
    """
    Exporta todos los productos existentes en la base de datos a un archivo Excel
    
    Args:
        session: Sesión de base de datos
        filename: Nombre del archivo a generar
    
    Returns:
        str: Ruta del archivo generado
    """
    
    # Crear workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    
    # Definir encabezados (mismo formato que el template)
    headers = [
        "serial_number",
        "name", 
        "description",
        "brand_id",
        "category_id",
        "warranty_time",
        "warranty_unit",
        "cost",
        "retail_price",
        "status",
        "variant_color",
        "variant_size",
        "variant_size_unit",
        "variant_unit",
        "variant_branch_id",
        "variant_stock",
        "variant_min_stock",
        "image_paths"
    ]
    
    descriptions = [
        "Número de serie único del producto (requerido)",
        "Nombre del producto (requerido)",
        "Descripción detallada (opcional)",
        "ID de la marca (opcional, número)",
        "ID de la categoría (opcional, número)",
        "Tiempo de garantía (opcional, número)",
        "Unidad de garantía: DAYS, MONTHS, YEARS (opcional)",
        "Costo del producto (requerido, número decimal)",
        "Precio de venta (requerido, número decimal)",
        "Estado: ACTIVE, INACTIVE, DISCONTINUED (opcional, default: ACTIVE)",
        "Color de la variante: ROJO, AZUL, VERDE, AMARILLO, NARANJA, VIOLETA, ROSADO, MARRON, GRIS, BLANCO, NEGRO, BORDO (opcional)",
        "Talla/Tamaño de la variante (opcional, texto)",
        "Tipo de tamaño: CLOTHING, DIMENSIONS, WEIGHT, OTHER (opcional)",
        "Unidad: KG, G, LB, CM, M, INCH, XS, S, L, XL, XXL (opcional)",
        "ID de la sucursal (requerido para variante, número)",
        "Stock disponible (opcional, número, default: 0)",
        "Stock mínimo (opcional, número, default: 5)",
        "Rutas de imágenes separadas por ; (ejemplo: C:\\imgs\\img1.jpg;C:\\imgs\\img2.jpg)"
    ]
    
    # Escribir encabezados
    ws.append(headers)
    
    # Estilo para encabezados
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Escribir descripciones
    ws.append(descriptions)
    
    # Estilo para descripciones
    desc_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    desc_font = Font(italic=True, size=9)
    
    for cell in ws[2]:
        cell.fill = desc_fill
        cell.font = desc_font
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    # Obtener todos los productos con sus variantes
    statement = select(Product).options(
        selectinload(Product.variants).selectinload(ProductVariant.images)
    )
    products = session.exec(statement).all()
    
    # Exportar cada producto con sus variantes
    for product in products:
        if product.variants and len(product.variants) > 0:
            # Si tiene variantes, crear una fila por cada variante
            for variant in product.variants:
                # Procesar imágenes
                image_paths = ""
                if variant.images and len(variant.images) > 0:
                    # Las imágenes están como data URLs, indicar que son imágenes ya cargadas
                    image_paths = f"[{len(variant.images)} imágenes ya cargadas en BD]"
                
                row = [
                    product.serial_number,
                    product.name,
                    product.description or "",
                    product.brand_id or "",
                    product.category_id or "",
                    product.warranty_time or "",
                    product.warranty_unit.value if product.warranty_unit else "",
                    product.cost,
                    product.retail_price,
                    product.status.value,
                    variant.color.value if variant.color else "",
                    variant.size or "",
                    variant.size_unit.value if variant.size_unit else "",
                    variant.unit.value if variant.unit else "",
                    variant.branch_id or "",
                    variant.stock,
                    variant.min_stock,
                    image_paths
                ]
                ws.append(row)
        else:
            # Si no tiene variantes, crear una fila solo con datos del producto
            row = [
                product.serial_number,
                product.name,
                product.description or "",
                product.brand_id or "",
                product.category_id or "",
                product.warranty_time or "",
                product.warranty_unit.value if product.warranty_unit else "",
                product.cost,
                product.retail_price,
                product.status.value,
                "",  # variant_color
                "",  # variant_size
                "",  # variant_size_unit
                "",  # variant_unit
                "",  # variant_branch_id
                "",  # variant_stock
                "",  # variant_min_stock
                ""   # image_paths
            ]
            ws.append(row)
    
    # Ajustar anchos de columnas
    column_widths = {
        'A': 20,  # serial_number
        'B': 30,  # name
        'C': 40,  # description
        'D': 12,  # brand_id
        'E': 12,  # category_id
        'F': 15,  # warranty_time
        'G': 15,  # warranty_unit
        'H': 12,  # cost
        'I': 12,  # retail_price
        'J': 15,  # status
        'K': 15,  # variant_color
        'L': 15,  # variant_size
        'M': 18,  # variant_size_unit
        'N': 15,  # variant_unit
        'O': 18,  # variant_branch_id
        'P': 15,  # variant_stock
        'Q': 18,  # variant_min_stock
        'R': 50,  # image_paths
    }
    
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    # Establecer altura de filas
    ws.row_dimensions[1].height = 40
    ws.row_dimensions[2].height = 60
    
    # Crear hoja de referencia
    ws_ref = wb.create_sheet("Referencia")
    
    reference_data = {
        "Unidades de Garantía": ["DAYS", "MONTHS", "YEARS"],
        "Estados de Producto": ["ACTIVE", "INACTIVE", "DISCONTINUED"],
        "Colores": ["ROJO", "AZUL", "VERDE", "AMARILLO", "NARANJA", "VIOLETA", 
                   "ROSADO", "MARRON", "GRIS", "BLANCO", "NEGRO", "BORDO"],
        "Unidades de Tamaño": ["CLOTHING", "DIMENSIONS", "WEIGHT", "OTHER"],
        "Unidades": ["KG", "G", "LB", "CM", "M", "INCH", "XS", "S", "L", "XL", "XXL"]
    }
    
    col_idx = 1
    for title, values in reference_data.items():
        ws_ref.cell(row=1, column=col_idx, value=title)
        ws_ref.cell(row=1, column=col_idx).font = Font(bold=True)
        ws_ref.cell(row=1, column=col_idx).fill = PatternFill(
            start_color="70AD47", end_color="70AD47", fill_type="solid"
        )
        ws_ref.cell(row=1, column=col_idx).font = Font(color="FFFFFF", bold=True)
        
        for idx, value in enumerate(values, start=2):
            ws_ref.cell(row=idx, column=col_idx, value=value)
        
        ws_ref.column_dimensions[ws_ref.cell(row=1, column=col_idx).column_letter].width = 25
        col_idx += 1
    
    # Guardar archivo
    wb.save(filename)
    
    return filename
